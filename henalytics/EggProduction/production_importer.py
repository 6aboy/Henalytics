import csv
from datetime import date, datetime
from decimal import Decimal
from io import StringIO

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from .models import ProductionLog


PRODUCTION_IMPORT_SESSION_KEY = 'production_log_import_preview'


HEADER_ALIASES = {
    'age': {'age week', 'age weeks', 'age_week', 'age'},
    'date': {'date', 'log date', 'log_date'},
    'bird_no': {'bird no.', 'bird no', 'bird number', 'bird_no', 'hen count', 'hens'},
    'dead_count': {'dead', 'dead count', 'dead_count', 'mortality'},
    'culled_count': {'cull', 'culled', 'culled count', 'culled_count'},
    'eggs_total': {
        'eggs collected (pcs)',
        'eggs collected',
        'eggs total',
        'eggs_total',
        'pieces',
        'total eggs collected',
    },
    'pct_hen_day': {'%hd', '% hd', '% hen day', '% hen-day', 'pct hen day'},
    'pct_hen_housed': {'%hh', '% hh', '% hen housed', '% hen-housed', 'pct hen housed'},
    'feed_bags': {'feed (bags)', 'feed bags', 'feed_bags', 'feed'},
    'fcr': {'fcr/ kg eggs', 'fcr/kg eggs', 'fcr', 'fcr kg eggs'},
}


def build_preview(uploaded_file, flock, date_mode, conflict_strategy):
    rows = parse_production_file(uploaded_file, flock, date_mode)
    valid_rows = [row for row in rows if not row['errors']]
    valid_dates = [date.fromisoformat(row['log_date']) for row in valid_rows]
    existing_dates = set()
    replace_existing_count = 0

    if valid_dates:
        existing_dates = set(
            ProductionLog.objects.filter(
                flock=flock,
                log_date__in=valid_dates,
            ).values_list('log_date', flat=True)
        )
        if conflict_strategy == 'replace_range':
            replace_existing_count = ProductionLog.objects.filter(
                flock=flock,
                log_date__gte=min(valid_dates),
                log_date__lte=max(valid_dates),
            ).count()

    for row in rows:
        if row['errors']:
            row['status'] = 'Invalid'
        elif date.fromisoformat(row['log_date']) in existing_dates:
            row['status'] = 'Will update' if conflict_strategy == 'update' else 'Will skip'
        else:
            row['status'] = 'New'

    if conflict_strategy == 'replace_range':
        for row in rows:
            if not row['errors']:
                row['status'] = 'Will replace range'

    stats = {
        'total_rows': len(rows),
        'valid_rows': len(valid_rows),
        'invalid_rows': len(rows) - len(valid_rows),
        'new_rows': sum(1 for row in rows if row['status'] == 'New'),
        'existing_rows': sum(1 for row in rows if row['status'] in ('Will update', 'Will skip')),
        'replace_existing_count': replace_existing_count,
        'date_from': min(valid_dates).isoformat() if valid_dates else '',
        'date_to': max(valid_dates).isoformat() if valid_dates else '',
    }
    return {
        'file_name': uploaded_file.name,
        'flock_id': flock.id,
        'flock_label': str(flock),
        'date_mode': date_mode,
        'conflict_strategy': conflict_strategy,
        'stats': stats,
        'rows': rows,
    }


def import_preview(preview, user):
    rows = [row for row in preview['rows'] if not row['errors']]
    if not rows:
        return {'created': 0, 'updated': 0, 'skipped': 0, 'deleted': 0}

    flock_id = preview['flock_id']
    conflict_strategy = preview['conflict_strategy']
    dates = [date.fromisoformat(row['log_date']) for row in rows]
    created = 0
    updated = 0
    skipped = 0
    deleted = 0

    with transaction.atomic():
        if conflict_strategy == 'replace_range':
            deleted, _ = ProductionLog.objects.filter(
                flock_id=flock_id,
                log_date__gte=min(dates),
                log_date__lte=max(dates),
            ).delete()

        for row in rows:
            log_date = date.fromisoformat(row['log_date'])
            existing = ProductionLog.objects.filter(flock_id=flock_id, log_date=log_date).first()
            if existing and conflict_strategy == 'skip':
                skipped += 1
                continue

            defaults = {
                'age_weeks': 0,
                'age_days': 0,
                'hen_count': row['bird_no'] or 0,
                'dead_count': row['dead_count'] or 0,
                'culled_count': row['culled_count'] or 0,
                'feed_bags': row['feed_bags'] or 0,
                'eggs_total': row['eggs_total'],
                'pct_hen_day': Decimal('0.00'),
                'pct_hen_housed': Decimal('0.00'),
                'remarks': f"Imported from {preview['file_name']}.",
                'entered_by': user,
            }
            _, was_created = ProductionLog.objects.update_or_create(
                flock_id=flock_id,
                log_date=log_date,
                defaults=defaults,
            )
            if was_created:
                created += 1
            else:
                updated += 1

        ProductionLog.recalculate_flock_snapshots(flock_id)

    return {'created': created, 'updated': updated, 'skipped': skipped, 'deleted': deleted}


def parse_production_file(uploaded_file, flock, date_mode):
    filename = uploaded_file.name.lower()
    if filename.endswith('.csv'):
        raw_rows = _csv_rows(uploaded_file)
    else:
        raw_rows = _excel_rows(uploaded_file)

    header_index, columns = _find_header(raw_rows)
    if header_index is None:
        return [{
            'row_number': 0,
            'log_date': '',
            'display_date': '',
            'bird_no': None,
            'dead_count': 0,
            'culled_count': 0,
            'feed_bags': 0,
            'eggs_total': 0,
            'status': 'Invalid',
            'errors': ['Could not find the production headers.'],
        }]

    rows = []
    current_date = None
    seen_dates = set()
    today = timezone.localdate()

    for row_number, raw_row in enumerate(raw_rows[header_index + 1:], start=header_index + 2):
        raw_date = _value(raw_row, columns.get('date'))
        eggs_total = _int_or_none(_value(raw_row, columns.get('eggs_total')))
        bird_no = _int_or_none(_value(raw_row, columns.get('bird_no')))
        if raw_date in (None, '') and eggs_total is None and bird_no is None:
            continue

        errors = []
        log_date = _resolve_date(raw_date, current_date, date_mode)
        if log_date is None:
            errors.append('Date could not be read.')
        else:
            current_date = log_date
            if log_date < flock.date_started:
                errors.append('Date is before the selected flock start date.')
            if log_date > today:
                errors.append('Date is in the future.')
            if log_date in seen_dates:
                errors.append('Duplicate date inside the uploaded file.')
            seen_dates.add(log_date)

        if eggs_total is None:
            errors.append('Eggs collected is missing.')
        elif eggs_total < 0:
            errors.append('Eggs collected cannot be negative.')

        rows.append({
            'row_number': row_number,
            'log_date': log_date.isoformat() if log_date else '',
            'display_date': log_date.strftime('%b %d, %Y') if log_date else '',
            'bird_no': bird_no,
            'dead_count': _int_or_zero(_value(raw_row, columns.get('dead_count'))),
            'culled_count': _int_or_zero(_value(raw_row, columns.get('culled_count'))),
            'feed_bags': _int_or_zero(_value(raw_row, columns.get('feed_bags'))),
            'eggs_total': eggs_total or 0,
            'status': '',
            'errors': errors,
        })

    return rows


def _excel_rows(uploaded_file):
    uploaded_file.seek(0)
    workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    return [list(row) for row in worksheet.iter_rows(values_only=True)]


def _csv_rows(uploaded_file):
    uploaded_file.seek(0)
    text = uploaded_file.read().decode('utf-8-sig')
    return [row for row in csv.reader(StringIO(text))]


def _find_header(rows):
    for index, row in enumerate(rows[:12]):
        normalized = [_normalize(value) for value in row]
        columns = {}
        for field, aliases in HEADER_ALIASES.items():
            for column_index, header in enumerate(normalized):
                if header in aliases:
                    columns[field] = column_index
                    break
        if 'date' in columns and 'eggs_total' in columns:
            return index, columns
    return None, {}


def _resolve_date(value, previous_date, date_mode):
    full_date = _full_date(value, date_mode)
    if full_date:
        return full_date

    day = _int_or_none(value)
    if day is None or previous_date is None or not 1 <= day <= 31:
        return None

    year = previous_date.year
    month = previous_date.month
    candidate = _safe_date(year, month, day)
    while candidate is None or (candidate <= previous_date and day < previous_date.day):
        month += 1
        if month > 12:
            month = 1
            year += 1
        candidate = _safe_date(year, month, day)
    return candidate


def _full_date(value, date_mode):
    if isinstance(value, datetime):
        raw_date = value.date()
    elif isinstance(value, date):
        raw_date = value
    elif isinstance(value, str):
        raw_date = _parse_date_string(value)
    else:
        return None

    if not raw_date:
        return None
    if date_mode == 'clsu_markers':
        return _safe_date(raw_date.year, raw_date.day, raw_date.month)
    return raw_date


def _parse_date_string(value):
    value = str(value).strip()
    if not value:
        return None
    for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%m-%d-%Y', '%m/%d/%y', '%m-%d-%y'):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


def _safe_date(year, month, day):
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _value(row, index):
    if index is None or index >= len(row):
        return None
    return row[index]


def _normalize(value):
    return ' '.join(str(value or '').strip().lower().replace('_', ' ').split())


def _int_or_none(value):
    if value in (None, ''):
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _int_or_zero(value):
    return _int_or_none(value) or 0
