import csv
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from io import StringIO

from django.db import transaction
from openpyxl import load_workbook

from .models import SalesItem, SalesTransaction


SALES_IMPORT_SESSION_KEY = 'sales_import_preview'


GRADE_LABELS = {
    'jumbo': 'Jumbo',
    'xl': 'XL',
    'large': 'Large',
    'medium': 'Medium',
    'small': 'Small',
    'pullets': 'Pullets',
    'pewee': 'Pewee',
    'broken': 'Broken',
    'cull': 'Cull',
    'sack': 'Sack',
}


GRADE_ALIASES = {
    'jumbo': 'jumbo',
    'xl': 'xl',
    'x large': 'xl',
    'x-large': 'xl',
    'extra large': 'xl',
    'large': 'large',
    'medium': 'medium',
    'small': 'small',
    'pullets': 'pullets',
    'pullet': 'pullets',
    'pewee': 'pewee',
    'peewee': 'pewee',
    'broken': 'broken',
    'cracked': 'broken',
    'cull': 'cull',
    'sack': 'sack',
}


def build_sales_preview(uploaded_file, date_mode, conflict_strategy):
    rows = parse_sales_file(uploaded_file, date_mode)
    valid_rows = [row for row in rows if not row['errors']]
    valid_dates = [date.fromisoformat(row['sale_date']) for row in valid_rows]
    existing_keys = set()
    replace_existing_count = 0

    if valid_rows:
        existing_keys = set(
            SalesTransaction.objects.filter(
                sale_date__in=valid_dates,
                or_number__in=[row['or_number'] for row in valid_rows],
            ).values_list('sale_date', 'or_number')
        )
        if conflict_strategy == 'replace_range':
            replace_existing_count = SalesTransaction.objects.filter(
                sale_date__gte=min(valid_dates),
                sale_date__lte=max(valid_dates),
            ).count()

    for row in rows:
        if row['errors']:
            row['status'] = 'Invalid'
            continue
        row_date = date.fromisoformat(row['sale_date'])
        if conflict_strategy == 'replace_range':
            row['status'] = 'Will replace range'
        elif (row_date, row['or_number']) in existing_keys:
            row['status'] = 'Will update' if conflict_strategy == 'update' else 'Will skip'
        else:
            row['status'] = 'New'

    stats = {
        'total_rows': len(rows),
        'valid_rows': len(valid_rows),
        'invalid_rows': len(rows) - len(valid_rows),
        'new_rows': sum(1 for row in rows if row['status'] == 'New'),
        'existing_rows': sum(1 for row in rows if row['status'] in ('Will update', 'Will skip')),
        'replace_existing_count': replace_existing_count,
        'date_from': min(valid_dates).isoformat() if valid_dates else '',
        'date_to': max(valid_dates).isoformat() if valid_dates else '',
        'total_amount': str(sum(Decimal(row['total_amount']) for row in valid_rows)),
        'total_pieces': sum(row['total_pieces'] for row in valid_rows),
    }
    return {
        'file_name': uploaded_file.name,
        'date_mode': date_mode,
        'conflict_strategy': conflict_strategy,
        'stats': stats,
        'rows': rows,
    }


def import_sales_preview(preview, user):
    rows = [row for row in preview['rows'] if not row['errors']]
    if not rows:
        return {'created': 0, 'updated': 0, 'skipped': 0, 'deleted': 0}

    conflict_strategy = preview['conflict_strategy']
    dates = [date.fromisoformat(row['sale_date']) for row in rows]
    created = 0
    updated = 0
    skipped = 0
    deleted = 0

    with transaction.atomic():
        if conflict_strategy == 'replace_range':
            deleted, _ = SalesTransaction.objects.filter(
                sale_date__gte=min(dates),
                sale_date__lte=max(dates),
            ).delete()

        for row in rows:
            sale_date = date.fromisoformat(row['sale_date'])
            transaction_obj = SalesTransaction.objects.filter(
                sale_date=sale_date,
                or_number=row['or_number'],
            ).first()
            if transaction_obj and conflict_strategy == 'skip':
                skipped += 1
                continue

            if transaction_obj:
                transaction_obj.recorded_by = user
                transaction_obj.notes = f"Imported from {preview['file_name']}."
                transaction_obj.save(update_fields=['recorded_by', 'notes', 'updated_at'])
                transaction_obj.items.all().delete()
                updated += 1
            else:
                transaction_obj = SalesTransaction.objects.create(
                    sale_date=sale_date,
                    or_number=row['or_number'],
                    recorded_by=user,
                    notes=f"Imported from {preview['file_name']}.",
                )
                created += 1

            for item in row['items']:
                SalesItem.objects.create(
                    transaction=transaction_obj,
                    grade=item['grade'],
                    quantity_pieces=item['quantity_pieces'],
                    amount=Decimal(item['amount']),
                )

    return {'created': created, 'updated': updated, 'skipped': skipped, 'deleted': deleted}


def parse_sales_file(uploaded_file, date_mode):
    if uploaded_file.name.lower().endswith('.csv'):
        return _parse_flat_csv(uploaded_file, date_mode)
    return _parse_sales_workbook(uploaded_file, date_mode)


def _parse_sales_workbook(uploaded_file, date_mode):
    uploaded_file.seek(0)
    workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    parsed_rows = []
    seen_keys = set()

    for sheet_name in workbook.sheetnames:
        rows = [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
        category_pairs = _category_pairs(rows)
        current_date = None

        for row_number, row in enumerate(rows, start=1):
            first_value = _value(row, 0)
            marker_date = _resolve_date(first_value, date_mode)
            if marker_date and _value(row, 1) in (None, ''):
                current_date = marker_date
                continue

            or_number = _or_number(_value(row, 1))
            if not or_number or not current_date:
                continue

            items = []
            for grade, pcs_column, amount_column in category_pairs:
                quantity = _int_or_zero(_value(row, pcs_column))
                amount = _decimal_or_zero(_value(row, amount_column))
                if quantity <= 0 and amount <= 0:
                    continue
                items.append({
                    'grade': grade,
                    'grade_label': GRADE_LABELS.get(grade, grade.title()),
                    'quantity_pieces': quantity,
                    'amount': str(amount),
                })

            if not items:
                continue

            parsed_rows.append(_sales_row(
                sheet_name,
                row_number,
                current_date,
                or_number,
                items,
                seen_keys,
            ))

    return parsed_rows


def _parse_flat_csv(uploaded_file, date_mode):
    uploaded_file.seek(0)
    text = uploaded_file.read().decode('utf-8-sig')
    reader = csv.DictReader(StringIO(text))
    grouped = defaultdict(list)
    row_numbers = {}
    seen_keys = set()

    for row_number, row in enumerate(reader, start=2):
        sale_date = _resolve_date(row.get('sale_date') or row.get('date'), date_mode)
        or_number = _or_number(row.get('or_number') or row.get('or') or row.get('OR'))
        grade = _normalize_grade(row.get('grade') or row.get('size') or row.get('egg_size'))
        quantity = _int_or_zero(row.get('quantity_pieces') or row.get('pcs') or row.get('pieces'))
        amount = _decimal_or_zero(row.get('amount') or row.get('total_amount'))
        if not sale_date or not or_number:
            key = (row_number, row_number)
        else:
            key = (sale_date.isoformat(), or_number)
        row_numbers.setdefault(key, row_number)
        grouped[key].append({
            'sale_date': sale_date,
            'or_number': or_number,
            'grade': grade,
            'grade_label': GRADE_LABELS.get(grade, str(grade or '').title()),
            'quantity_pieces': quantity,
            'amount': str(amount),
        })

    parsed_rows = []
    for key, items in grouped.items():
        first = items[0]
        sale_date = first.get('sale_date')
        or_number = first.get('or_number')
        valid_items = [
            {
                'grade': item['grade'],
                'grade_label': item['grade_label'],
                'quantity_pieces': item['quantity_pieces'],
                'amount': item['amount'],
            }
            for item in items
            if item['grade'] and (item['quantity_pieces'] > 0 or Decimal(item['amount']) > 0)
        ]
        parsed_rows.append(_sales_row('CSV', row_numbers[key], sale_date, or_number, valid_items, seen_keys))

    return parsed_rows


def _sales_row(sheet_name, row_number, sale_date, or_number, items, seen_keys):
    errors = []
    if not sale_date:
        errors.append('Date could not be read.')
    if not or_number:
        errors.append('OR number is missing.')
    if not items:
        errors.append('No sales items found.')

    key = (sale_date, or_number)
    if sale_date and or_number and key in seen_keys:
        errors.append('Duplicate OR/date inside the uploaded file.')
    if sale_date and or_number:
        seen_keys.add(key)

    return {
        'sheet_name': sheet_name,
        'row_number': row_number,
        'sale_date': sale_date.isoformat() if sale_date else '',
        'display_date': sale_date.strftime('%b %d, %Y') if sale_date else '',
        'or_number': or_number or '',
        'items': items,
        'items_label': ', '.join(item['grade_label'] for item in items) or '-',
        'total_pieces': sum(item['quantity_pieces'] for item in items),
        'total_amount': str(sum(Decimal(item['amount']) for item in items)),
        'status': '',
        'errors': errors,
    }


def _category_pairs(rows):
    if len(rows) < 2:
        return []
    labels = rows[1]
    pairs = []
    for index, label in enumerate(labels):
        grade = _normalize_grade(label)
        if grade and index + 1 < len(labels):
            pairs.append((grade, index, index + 1))
    return pairs


def _normalize_grade(value):
    normalized = ' '.join(str(value or '').strip().lower().replace('_', ' ').split())
    return GRADE_ALIASES.get(normalized)


def _resolve_date(value, date_mode):
    if isinstance(value, datetime):
        raw_date = value.date()
    elif isinstance(value, date):
        raw_date = value
    elif isinstance(value, str):
        raw_date = _parse_date_string(value)
    else:
        return None

    if not raw_date:
        return None
    if date_mode == 'clsu_markers':
        return _safe_date(raw_date.year, raw_date.day, raw_date.month)
    return raw_date


def _parse_date_string(value):
    value = str(value).strip()
    if not value:
        return None
    for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%m-%d-%Y', '%m/%d/%y', '%m-%d-%y'):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


def _safe_date(year, month, day):
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _value(row, index):
    if index is None or index >= len(row):
        return None
    return row[index]


def _or_number(value):
    if value in (None, ''):
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _int_or_none(value):
    if value in (None, ''):
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _int_or_zero(value):
    return _int_or_none(value) or 0


def _decimal_or_zero(value):
    if value in (None, ''):
        return Decimal('0.00')
    try:
        return Decimal(str(value)).quantize(Decimal('0.01'))
    except Exception:
        return Decimal('0.00')
